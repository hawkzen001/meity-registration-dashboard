"""
Fetch latest 'Main Data' directly from Google Sheets CSV export,
cross-reference against IDEMI Animation Department existing enrollment files
(last 3 academic years), and generate src/data/registrations.json.

Excludes candidates without documents.
Marks already enrolled candidates with already_enrolled_at_idemi = True / False.
"""

import urllib.request
import csv
import json
import os
import io
import sys
import re
import openpyxl
from datetime import datetime
from pathlib import Path

os.makedirs('src/data', exist_ok=True)

# ── Google Sheet identifiers ────────────────────────────────────────────────
SHEET_ID = "1vL2_xcWv_IvgqGIGrtWIQFCXXQ9qHAvuTpfdY_uKFyM"
GID      = "2069006072"    # gid of "Main Data" tab
CSV_URL  = (
    f"https://docs.google.com/spreadsheets/d/{SHEET_ID}"
    f"/export?format=csv&gid={GID}"
)

# ── Course seat capacities ───────────────────────────────────────────────────
COURSE_CAPACITIES = {
    'Graphics & Web Designer Assistant (NSQF Level 4)':       {'total':180,'SC':27,'ST':14,'OBC':48,'EWS':18,'General':73},
    'Multimedia and Animation Associate (NSQF Level 4)':      {'total':120,'SC':18,'ST': 9,'OBC':32,'EWS':12,'General':49},
    'VFX Associate (NSQF Level 4)':                           {'total': 90,'SC':13,'ST': 7,'OBC':24,'EWS': 9,'General':37},
    '3D Printing':                                             {'total': 60,'SC': 9,'ST': 5,'OBC':16,'EWS': 6,'General':24},
    '3D Animator Assistant (NSQF Level 3)':                   {'total': 60,'SC': 9,'ST': 5,'OBC':16,'EWS': 6,'General':24},
    'Jr. Web Designer (NSQF Level 4)':                        {'total': 60,'SC': 9,'ST': 5,'OBC':16,'EWS': 6,'General':24},
    'Sr. Technician -Mechatronics (NSQF Level 4.5)':          {'total': 30,'SC': 5,'ST': 2,'OBC': 8,'EWS': 3,'General':12},
    'Film Compositor Assistant (NSQF Level 4)':               {'total': 30,'SC': 5,'ST': 2,'OBC': 8,'EWS': 3,'General':12},
}

# ── Document-upload column indices (1-based) ─────────────────────────────────
DOC_COL_MAP = {
    "ssc":        [34, 48, 58],
    "hsc":        [35, 59, 69],
    "graduation": [36, 50, 60],
    "age_proof":  [37, 51, 61, 70],
    "photo":      [38, 52, 62, 68],
    "identity":   [39, 53, 63, 71],
    "caste":      [44, 54, 64, 75, 83],
    "ews_income": [45],
    "diploma":    [49],
}

# ── Normalise functions ──────────────────────────────────────────────────────
def norm(s):
    if not s: return ""
    s = str(s).lower().strip()
    s = re.sub(r'\b(mr|mrs|ms|dr|shri|smt)\.?\s*', '', s)
    s = re.sub(r"[^a-z0-9\s]", " ", s)
    return re.sub(r"\s+", " ", s).strip()

def norm_gender(v):
    l = v.lower()
    if "female" in l or "women" in l or "woman" in l: return "Female"
    if "male"   in l or "boy"   in l or "man"   in l: return "Male"
    if "trans"  in l: return "Transgender"
    return "Other" if v and v != "None" else "Male"

def norm_category(v):
    l = v.lower()
    if "scheduled caste" in l or "schedule caste" in l or " sc" in l or l.startswith("sc"): return "SC"
    if "scheduled tribe" in l or "schedule tribe" in l or " st" in l or l.startswith("st"): return "ST"
    if "obc" in l or "other backward" in l: return "OBC"
    if "ews" in l or "economically weaker" in l: return "EWS"
    if "women" in l or "female" in l: return "General (Women)"
    return "General"

def norm_qual(v):
    l = v.lower()
    if "degree" in l or "graduat" in l: return "Degree / Graduation"
    if "12th"   in l or "iti"    in l or "hsc" in l: return "12th / ITI"
    if "diploma" in l: return "Diploma"
    if "10th"   in l or "ssc"   in l: return "10th Pass"
    return v or "Other"

# ── 1. Load Existing Enrollment Data ─────────────────────────────────────────
ENROLL_DIR = Path("Exisiting Enrollment Data - Animation Department")
FILE_NAME_COLS = {
    "_Animation Batches 25-26.xlsx":   [4, 16, 27],
    "Ani - MeitY Batches 24-25.xlsx":  [4, 16, 27],
    "Animation Batches 26-27.xlsx":    [4, 15, 27, 38, 49],
}

enrolled_data = {}
if ENROLL_DIR.exists():
    for xlsx in sorted(ENROLL_DIR.glob("*.xlsx")):
        fname = xlsx.name
        cols = FILE_NAME_COLS.get(fname, [4])
        try:
            wb = openpyxl.load_workbook(xlsx, data_only=True)
            for sname in wb.sheetnames:
                ws = wb[sname]
                row1 = [cell.value for cell in ws[1]]
                for row in ws.iter_rows(min_row=3, values_only=True):
                    if not row: continue
                    for c in cols:
                        if len(row) >= c and row[c-1]:
                            val = str(row[c-1]).strip()
                            if val and not val.replace('.','').isdigit():
                                nname = norm(val)
                                if len(nname.split()) >= 2:
                                    # Find enrollment number: First non-blank cell before the Name column
                                    en_no = "N/A"
                                    for i in range(c - 2, max(-1, c - 6), -1):
                                        if i < len(row) and row[i] is not None and str(row[i]).strip() != "":
                                            en_no = str(row[i]).replace('.0', '').strip()
                                            break
                                    
                                    # Find course name: First non-blank string before or at Name column in Row 1
                                    c_name = "Unknown Course"
                                    for i in range(c - 1, -1, -1):
                                        if i < len(row1) and row1[i] and str(row1[i]).strip():
                                            c_name = str(row1[i]).strip()
                                            break
                                    
                                    enrolled_data[nname] = {
                                        "name": val,
                                        "course": c_name,
                                        "enrollment_no": en_no
                                    }
        except Exception as e:
            print(f"Warning loading {fname}: {e}")

print(f"Extracted {len(enrolled_data)} unique enrolled student names from IDEMI Animation files.")

def is_enrolled_match(name):
    nw = set(norm(name).split())
    if not nw: return False, ""
    for en, data in enrolled_data.items():
        enw = set(en.split())
        common = nw & enw
        shorter = min(len(nw), len(enw))
        if nw == enw:
            return True, f"exact: '{data['name']}' | {data['course']} | #{data['enrollment_no']}"
        if len(common) >= 2 and shorter > 0 and len(common) / shorter >= 0.75:
            return True, f"fuzzy: '{data['name']}' | {data['course']} | #{data['enrollment_no']}"
    return False, ""

# ── 2. Fetch CSV ─────────────────────────────────────────────────────────────
print("Fetching latest data from Google Sheets …")
try:
    req = urllib.request.Request(CSV_URL, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=30) as resp:
        raw_bytes = resp.read()
    raw_text = raw_bytes.decode("utf-8")
    print(f"  Downloaded {len(raw_bytes):,} bytes from Google Sheets OK")
except Exception as e:
    print(f"ERROR: Could not fetch Google Sheets data: {e}")
    sys.exit(1)

reader   = csv.reader(io.StringIO(raw_text))
all_rows = list(reader)
headers  = all_rows[0] if all_rows else []
data_rows = all_rows[1:]

print(f"  Sheet rows (incl. header): {len(all_rows)}  |  Data rows: {len(data_rows)}")

# ── 3. Process rows ─────────────────────────────────────────────────────────
def cell(row, col_1based, default=""):
    idx = col_1based - 1
    return row[idx].strip() if idx < len(row) and row[idx] else default

all_records          = []
course_stats         = {}
category_counts      = {}
gender_counts        = {}
qualification_counts = {}

already_enrolled_count = 0

for row in data_rows:
    timestamp   = cell(row, 1)
    email       = cell(row, 2)
    name        = cell(row, 3)
    father_name = cell(row, 4)
    gender_raw  = cell(row, 5)
    category_raw= cell(row, 6)
    dob         = cell(row, 7)
    qualification= cell(row, 8)
    address     = cell(row, 9)
    alt_phone   = cell(row, 10)
    whatsapp    = cell(row, 14)

    if not (name or email or timestamp):
        continue

    # clean phone
    for attr in ("whatsapp", "alt_phone"):
        v = locals()[attr]
        if v.endswith(".0"): locals().__setitem__(attr, v[:-2])
    whatsapp  = whatsapp.rstrip(".0") if whatsapp.endswith(".0")  else whatsapp
    alt_phone = alt_phone.rstrip(".0") if alt_phone.endswith(".0") else alt_phone

    # documents
    docs = {}
    for key, cols in DOC_COL_MAP.items():
        for c in cols:
            v = cell(row, c)
            if v.startswith("http"):
                docs[key] = v
                break

    # EXCLUDE candidates without any uploaded document
    if not docs:
        continue

    # category
    cat_raw = ""
    for c in [6, 76, 77, 78]:
        v = cell(row, c)
        if v: cat_raw = v; break

    # course
    course = "Other / Unspecified"
    for c in [11, 46, 47, 74]:
        v = cell(row, c)
        if v: course = v; break

    # reference
    reference = ""
    for c in [12, 42, 56, 66, 72]:
        v = cell(row, c)
        if v: reference = v; break

    std_gender  = norm_gender(gender_raw)
    std_cat     = norm_category(cat_raw)
    std_qual    = norm_qual(qualification)

    # Check enrollment status
    enrolled, match_detail = is_enrolled_match(name)
    if enrolled:
        already_enrolled_count += 1

    rec = {
        "id":                       len(all_records) + 1,
        "sheet":                    "Main Data",
        "timestamp":                timestamp[:19],
        "name":                     name.title() or "Unknown Applicant",
        "email":                    email,
        "father_name":              father_name.title(),
        "gender":                   std_gender,
        "category":                 std_cat,
        "raw_category":             cat_raw,
        "dob":                      dob[:10],
        "qualification":            std_qual,
        "address":                  address,
        "whatsapp":                 whatsapp,
        "alt_phone":                alt_phone,
        "course":                   course,
        "reference":                reference,
        "documents":                docs,
        "doc_count":                len(docs),
        "already_enrolled_at_idemi":enrolled,
        "enrollment_match_detail":  match_detail,
    }

    all_records.append(rec)
    gender_counts[std_gender]        = gender_counts.get(std_gender, 0) + 1
    category_counts[std_cat]         = category_counts.get(std_cat, 0) + 1
    course_stats[course]             = course_stats.get(course, 0) + 1
    qualification_counts[std_qual]   = qualification_counts.get(std_qual, 0) + 1

# ── 4. Write output ─────────────────────────────────────────────────────────
payload = {
    "generated_at":           datetime.now().strftime("%Y-%m-%dT%H:%M:%S"),
    "source":                 "Google Sheets Live Sync",
    "total_records":          len(all_records),
    "already_enrolled_count": already_enrolled_count,
    "not_enrolled_count":     len(all_records) - already_enrolled_count,
    "gender_counts":          gender_counts,
    "category_counts":        category_counts,
    "course_counts":          course_stats,
    "qualification_counts":   qualification_counts,
    "course_capacities":      COURSE_CAPACITIES,
    "records":                all_records,
}

with open("src/data/registrations.json", "w", encoding="utf-8") as f:
    json.dump(payload, f, indent=2, ensure_ascii=False)

print(f"\nDone! Exported {len(all_records):,} verified records to src/data/registrations.json")
print(f"  • Already Enrolled: {already_enrolled_count}")
print(f"  • Not Yet Enrolled (New Leads): {len(all_records) - already_enrolled_count}")
